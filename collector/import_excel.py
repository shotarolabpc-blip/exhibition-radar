"""既存の展示会イベントリスト（.xlsx / .csv）からの初期データ取込（初回のみ）。

使い方：
    python -m collector.import_excel "C:/path/展示会イベントリスト.xlsx" [--dry-run]

- 取り込む列：イベント名、開始日、終了日、会場、対応カテゴリ、概要、URL
- 取り込まない列：No、SA部アサイン、コムシスアサイン（その他の列も読まない）
- 入力ファイル自体はリポジトリに置かないこと（.gitignore で *.xlsx / 入力CSV を除外）
- 既存ExcelのURL一覧は config/known_urls.yaml に書き出し、known_urls 収集元の巡回対象にする
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import yaml

from collector.date_range import compute_period
from collector.exporter import append_run, archive_events, load_events, write_events
from collector.models import Event, compose_event_name, edition_id, is_same_edition, make_event_id
from collector.normalizer import (
    clean_text,
    normalize_categories,
    normalize_event_name,
    normalize_url,
    normalize_venue,
    parse_dates,
    split_event_name,
)
from collector.overrides import apply_overrides
from collector.region import resolve_region
from collector.settings import Config, load_config, load_overrides

JST = timezone(timedelta(hours=9))
IMPORT_COLUMNS = ("イベント名", "開始日", "終了日", "会場", "対応カテゴリ", "概要", "URL")
FORBIDDEN_COLUMNS = ("SA部アサイン", "コムシスアサイン")
SOURCE_NAME = "既存Excel取込"


def _pick_columns(header: list[Any]) -> dict[str, int]:
    names = [clean_text(h) for h in header]
    index = {name: names.index(name) for name in IMPORT_COLUMNS if name in names}
    if "イベント名" not in index:
        raise ValueError(f"見出し行に「イベント名」がありません：{names}")
    return index


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        return []
    cols = _pick_columns(rows[0])
    return [{name: (r[i] if i < len(r) else "") for name, i in cols.items()} for r in rows[1:]]


def read_xlsx_rows(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    """URL列はセルのハイパーリンク先を優先して読む（表示文字列だけのセル対策）。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows())
    header_idx = next(
        (i for i, row in enumerate(rows[:20]) if any(clean_text(c.value) == "イベント名" for c in row)),
        None,
    )
    if header_idx is None:
        raise ValueError("先頭20行に「イベント名」の見出しが見つかりません")
    cols = _pick_columns([c.value for c in rows[header_idx]])
    result: list[dict[str, Any]] = []
    for row in rows[header_idx + 1 :]:
        item: dict[str, Any] = {}
        for name, i in cols.items():
            cell = row[i] if i < len(row) else None
            value = cell.value if cell is not None else None
            if name == "URL" and cell is not None and cell.hyperlink is not None and cell.hyperlink.target:
                value = cell.hyperlink.target
            item[name] = value
        result.append(item)
    return result


def read_rows(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx_rows(path, sheet)
    return read_csv_rows(path)


def _iso(value: date | None) -> str:
    return value.isoformat() if value else ""


def _date_text(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date()
    return value


def row_to_event(row: dict[str, Any], cfg: Config, today: str) -> Event | None:
    name = normalize_event_name(row.get("イベント名"))
    if not name:
        return None
    parent, sub = split_event_name(name)

    ps = parse_dates(_date_text(row.get("開始日")))
    pe = parse_dates(_date_text(row.get("終了日")))
    start = ps.start
    end = pe.start or start
    if start and end and end < start:
        end = start

    review: list[str] = []
    year = str(ps.year) if ps.year else ""
    if not year:
        # 日付列に年が無い場合のみ、イベント名に明記された西暦を使う（推測はしない）
        m = re.search(r"(?<!\d)(20\d{2})(?!\d)", name)
        year = m.group(1) if m else ""
    if not year:
        review.append("開催年が不明")
    if not start:
        review.append("日程未定" if year else "日付なし")

    raw_url = clean_text(row.get("URL"))
    url = normalize_url(raw_url)
    if raw_url and not url:
        review.append(f"URL列がリンク文字列のみ（「{raw_url[:40]}」）")
    elif not url:
        review.append("URLなし")

    raw_venue = clean_text(row.get("会場"))
    venue = normalize_venue(raw_venue, cfg.venues)
    prefecture, region = resolve_region(venue, raw_venue, cfg.venues)
    if venue and not region:
        review.append("会場の地域が不明")

    needs_review = bool({"開催年が不明", "日付なし"} & set(review)) or any(r.startswith("URL列") for r in review)
    reason = "既存の展示会イベントリストから取込"
    if review:
        reason += "（" + "、".join(review) + "）"

    return Event(
        id=make_event_id(parent, sub, year),
        event_name=compose_event_name(parent, sub),
        parent_event_name=parent,
        sub_event_name=sub,
        start_date=_iso(start),
        end_date=_iso(end),
        date_status="fixed" if start else "unknown",
        venue=venue,
        prefecture=prefecture,
        region=region,
        categories=normalize_categories(row.get("対応カテゴリ"), cfg.categories),
        summary=clean_text(row.get("概要")),
        url=url,
        source=SOURCE_NAME,
        source_type="import",
        keywords=[],
        confidence=90,
        reason=reason,
        status="needs_review" if needs_review else "unchanged",
        first_seen=today,
        last_updated=today,
        last_checked=today,
        year=year,
        history=[{"date": today, "change": "既存Excelから取込"}],
    )


def build_known_urls(events: list[Event]) -> list[dict[str, str]]:
    seen: dict[str, dict[str, str]] = {}
    for e in events:
        if e.url and e.url not in seen:
            seen[e.url] = {"url": e.url, "name": e.parent_event_name}
    return sorted(seen.values(), key=lambda d: d["url"])


def write_known_urls(path: Path, urls: list[dict[str, str]]) -> None:
    header = "# 既知イベントURL（known_urls 収集元の巡回対象）\n# import_excel.py が既存リストのURLから生成。手で追加・削除してよい。\n"
    with path.open("w", encoding="utf-8", newline="\n") as f:
        f.write(header)
        yaml.safe_dump({"urls": urls}, f, allow_unicode=True, sort_keys=False, width=200)


def run_import(path: Path, cfg: Config, today: date, dry_run: bool = False, sheet: str | None = None) -> dict[str, Any]:
    today_s = today.isoformat()
    rows = read_rows(path, sheet)
    imported: dict[str, Event] = {}
    duplicates = 0
    all_events: list[Event] = []
    for row in rows:
        event = row_to_event(row, cfg, today_s)
        if event is None:
            continue
        all_events.append(event)
        other = imported.get(event.id)
        if other is not None and not is_same_edition(other, event):
            event.id = edition_id(event)
            other = imported.get(event.id)
        if other is not None:
            duplicates += 1
            continue
        imported[event.id] = event

    data_dir = cfg.docs_data_dir
    existing = {e.id: e for e in load_events(data_dir / "events.json")}
    added = [e for e in imported.values() if e.id not in existing]
    merged = list(existing.values()) + added
    merged = apply_overrides(merged, load_overrides(), cfg.venues, cfg.categories, today_s)
    archive_days = int(cfg.section("period").get("archive_after_days", 30))

    summary = {
        "rows": len(rows),
        "events": len(all_events),
        "unique": len(imported),
        "duplicates_in_file": duplicates,
        "added": len(added),
        "status": dict(Counter(e.status for e in merged)),
        "unknown_venue_region": sorted({e.venue for e in merged if e.venue and not e.region}),
    }
    if dry_run:
        limit = (today - timedelta(days=archive_days)).isoformat()
        summary["would_archive"] = sum(1 for e in merged if (e.end_date or e.start_date) and (e.end_date or e.start_date) < limit)
        return summary

    remaining = archive_events(merged, today, archive_days, cfg.archive_dir)
    write_events(remaining, data_dir)
    write_known_urls(cfg.root_dir / "config" / "known_urls.yaml", build_known_urls(all_events))
    period = compute_period(today, int(cfg.section("period").get("months_ahead", 5)))
    append_run(
        data_dir,
        {
            "run_at": datetime.now(JST).isoformat(timespec="seconds"),
            "type": "import",
            "period": period.to_dict(),
            "sources_total": 0,
            "sources_failed": 0,
            "pages_fetched": 0,
            "gemini_requests": 0,
            "new": 0,
            "updated": 0,
            "needs_review": sum(1 for e in remaining if e.status == "needs_review"),
            "imported": len(added),
            "archived": len(merged) - len(remaining),
            "total": len(remaining),
            "errors": [],
        },
        int(cfg.section("output").get("runs_keep", 100)),
    )
    summary["archived"] = len(merged) - len(remaining)
    summary["published"] = len(remaining)
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="既存の展示会イベントリストを events.json に取り込む")
    parser.add_argument("path", type=Path, help=".xlsx または .csv（リポジトリ外に置くこと）")
    parser.add_argument("--sheet", help="xlsxのシート名（省略時は先頭シート）")
    parser.add_argument("--dry-run", action="store_true", help="書き込まずに集計だけ表示")
    parser.add_argument("--today", help="基準日 YYYY-MM-DD（テスト用）")
    args = parser.parse_args(argv)

    cfg = load_config()
    if args.path.resolve().is_relative_to(cfg.root_dir):
        print("入力ファイルはリポジトリの外に置いてください（社内情報の混入防止）", file=sys.stderr)
        return 2
    today = date.fromisoformat(args.today) if args.today else datetime.now(JST).date()
    summary = run_import(args.path, cfg, today, dry_run=args.dry_run, sheet=args.sheet)
    for key, value in summary.items():
        print(f"{key}: {value}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
